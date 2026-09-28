import io

from docx import Document


def _make_resume_bytes() -> bytes:
    document = Document()
    document.add_heading("Test Student Resume", level=1)
    document.add_paragraph("Skills: Python, SQL, Excel, Statistics, Communication")
    document.add_paragraph("Projects:")
    document.add_paragraph("Made a machine learning project for predicting sales.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _create_profile(client, auth_headers, user_id):
    resp = client.post(
        "/api/students/profiles",
        json={"user_id": user_id, "degree": "B.Tech", "branch": "CSE"},
        headers=auth_headers,
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


def test_resume_upload_extracts_text_and_syncs_skills(client, auth_headers, admin_headers):
    user_id = client.get("/api/auth/me", headers=auth_headers).json()["id"]
    _create_profile(client, auth_headers, user_id)

    for name in ["Python", "Statistics"]:
        resp = client.post("/api/skills", json={"name": name, "category": "technical"}, headers=admin_headers)
        assert resp.status_code == 201

    resume_bytes = _make_resume_bytes()
    files = {
        "file": (
            "resume.docx",
            resume_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    upload_resp = client.post("/api/resumes/upload", files=files, headers=auth_headers)
    assert upload_resp.status_code == 201, upload_resp.text
    resume = upload_resp.json()
    assert resume["analysis_status"] == "completed"
    assert "Python" in resume["extracted_text"]

    my_skills_resp = client.get("/api/skills/me", headers=auth_headers)
    assert my_skills_resp.status_code == 200
    skill_names = {row["skill_id"] for row in my_skills_resp.json()}
    assert len(skill_names) >= 1  # Statistics/Python auto-synced from resume text

    improve_resp = client.post(
        f"/api/resumes/{resume['id']}/improve",
        json={"target_career": "Data Analyst"},
        headers=auth_headers,
    )
    assert improve_resp.status_code == 200
    assert "Weaknesses" in improve_resp.json()["improvement"]


def test_resume_rejects_unsupported_file_type(client, auth_headers):
    user_id = client.get("/api/auth/me", headers=auth_headers).json()["id"]
    _create_profile(client, auth_headers, user_id)

    files = {"file": ("notes.txt", b"just text", "text/plain")}
    resp = client.post("/api/resumes/upload", files=files, headers=auth_headers)
    assert resp.status_code == 400


def test_resume_owner_isolation(client, register_and_login):
    headers_a = register_and_login(email="studenta@example.com", name="Student A")
    headers_b = register_and_login(email="studentb@example.com", name="Student B")
    user_a_id = client.get("/api/auth/me", headers=headers_a).json()["id"]
    _create_profile(client, headers_a, user_a_id)

    resume_bytes = _make_resume_bytes()
    files = {
        "file": (
            "resume.docx",
            resume_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    upload_resp = client.post("/api/resumes/upload", files=files, headers=headers_a)
    resume_id = upload_resp.json()["id"]

    forbidden_resp = client.get(f"/api/resumes/{resume_id}", headers=headers_b)
    assert forbidden_resp.status_code == 403

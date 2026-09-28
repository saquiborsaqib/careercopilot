from pathlib import Path

from docx import Document


def extract_text_from_docx(file_path: str | Path) -> str:
    document = Document(str(file_path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())

    return "\n".join(parts).strip()
